"""
Unit Tests for Supermarket Sales Intelligence Reporting Engine
===============================================================
Tests executive PDF document generation (ReportLab) and multi-sheet
Excel workbook creation (openpyxl).
"""

import io
from pathlib import Path
import unittest
import pandas as pd
import openpyxl

from src.reports import (
    generate_pdf_report,
    generate_excel_workbook,
    generate_csv_export,
)
from src.inventory import calculate_inventory_intelligence
from src.anomaly import detect_anomalies
from src.insights import generate_structured_business_insights
from src.database import (
    init_database,
    seed_database_if_empty,
    get_products_df,
    create_product,
    record_stock_movement,
    get_all_products,
    log_report,
    log_report_generation,
    normalize_report_type,
    get_reports_df,
    get_report_history,
    get_inventory_movements_df,
    get_database_stats,
    get_db_connection,
)


class TestReportingEngine(unittest.TestCase):
    """Test suite for PDF and Excel business report generation."""

    def setUp(self):
        """Create sample sales, inventory, insights, and anomaly fixtures."""
        self.sample_df = pd.DataFrame([
            {
                "Invoice ID": "INV-1001",
                "Branch": "A",
                "City": "Yangon",
                "Customer Type": "Member",
                "Gender": "Female",
                "Product": "Organic Apples",
                "Product Category": "Food and beverages",
                "Unit Price": 10.0,
                "Quantity": 5,
                "Sales": 50.0,
                "Date": "2023-01-15",
                "Payment": "Ewallet",
                "Rating": 4.5,
            },
            {
                "Invoice ID": "INV-1002",
                "Branch": "B",
                "City": "Mandalay",
                "Customer Type": "Normal",
                "Gender": "Male",
                "Product": "Smartphone Case",
                "Product Category": "Electronic accessories",
                "Unit Price": 20.0,
                "Quantity": 2,
                "Sales": 40.0,
                "Date": "2023-01-16",
                "Payment": "Cash",
                "Rating": 3.8,
            },
            {
                "Invoice ID": "INV-1003",
                "Branch": "C",
                "City": "Naypyitaw",
                "Customer Type": "Member",
                "Gender": "Female",
                "Product": "Face Wash",
                "Product Category": "Health and beauty",
                "Unit Price": 15.0,
                "Quantity": 4,
                "Sales": 60.0,
                "Date": "2023-01-17",
                "Payment": "Credit card",
                "Rating": 4.9,
            },
        ])

        # Generate anomalies & insights
        self.anomalies = detect_anomalies(self.sample_df)
        self.insights = generate_structured_business_insights(self.sample_df)

        # Sample inventory DataFrame
        self.inventory_df = pd.DataFrame([
            {
                "product_id": 1,
                "product_name": "Organic Apples",
                "category_name": "Food and beverages",
                "current_stock": 10,
                "safety_stock_level": 20,
                "unit_cost": 7.0,
                "unit_selling_price": 10.0,
                "daily_sales_velocity": 2.5,
                "days_of_stock_remaining": 4.0,
                "velocity_tier": "Fast Moving",
                "health_status": "Low Stock",
                "reorder_flag": 1,
                "recommended_reorder_qty": 50,
                "estimated_reorder_cost": 350.0,
            },
            {
                "product_id": 2,
                "product_name": "Smartphone Case",
                "category_name": "Electronic accessories",
                "current_stock": 45,
                "safety_stock_level": 15,
                "unit_cost": 12.0,
                "unit_selling_price": 20.0,
                "daily_sales_velocity": 0.8,
                "days_of_stock_remaining": 56.2,
                "velocity_tier": "Normal Moving",
                "health_status": "Healthy",
                "reorder_flag": 0,
                "recommended_reorder_qty": 0,
                "estimated_reorder_cost": 0.0,
            },
        ])

    def test_pdf_report_generation(self):
        """Verify PDF generation returns valid binary data starting with %PDF."""
        filter_summary = {
            "date_range": "2023-01-15 to 2023-01-17",
            "branch": ["A", "B", "C"],
            "city": ["Yangon", "Mandalay", "Naypyitaw"],
        }
        pdf_buf = generate_pdf_report(
            df=self.sample_df,
            dataset_name="Supermarket Baseline Fixture",
            filter_summary=filter_summary,
            insights=self.insights,
            anomalies=self.anomalies,
            inventory_df=self.inventory_df,
        )

        self.assertIsInstance(pdf_buf, io.BytesIO)
        pdf_bytes = pdf_buf.getvalue()
        self.assertGreater(len(pdf_bytes), 1000, "PDF should contain compiled pages.")
        self.assertTrue(pdf_bytes.startswith(b"%PDF"), "PDF binary header must start with %PDF.")

    def test_pdf_report_empty_dataframe(self):
        """Verify PDF report generates cleanly even with an empty DataFrame."""
        empty_df = pd.DataFrame(columns=["Invoice ID", "Branch", "City", "Sales", "Quantity", "Rating"])
        pdf_buf = generate_pdf_report(
            df=empty_df,
            dataset_name="Empty Scope Dataset",
            filter_summary={},
            insights=[],
            anomalies=[],
            inventory_df=None,
        )

        self.assertIsInstance(pdf_buf, io.BytesIO)
        pdf_bytes = pdf_buf.getvalue()
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))

    def test_excel_workbook_generation(self):
        """Verify multi-sheet Excel generation and validate sheet names and content."""
        filter_summary = {"branch": ["A", "B"]}
        xlsx_buf = generate_excel_workbook(
            df=self.sample_df,
            dataset_name="Test Store Scope",
            filter_summary=filter_summary,
            insights=self.insights,
            anomalies=self.anomalies,
            inventory_df=self.inventory_df,
        )

        self.assertIsInstance(xlsx_buf, io.BytesIO)
        xlsx_bytes = xlsx_buf.getvalue()
        self.assertGreater(len(xlsx_bytes), 2000, "Excel workbook should contain formatted tabs.")

        # Reopen workbook with openpyxl to inspect sheet structure
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        sheet_names = wb.sheetnames

        expected_sheets = [
            "Executive Summary",
            "Sales Transactions",
            "Product Performance",
            "Category Performance",
            "Inventory Intelligence",
            "Anomaly Audit Log",
            "Business Insights Matrix",
        ]
        for sheet in expected_sheets:
            self.assertIn(sheet, sheet_names, f"Expected sheet '{sheet}' missing from workbook.")

        # Check Executive Summary content
        ws_sum = wb["Executive Summary"]
        self.assertEqual(ws_sum["A1"].value, "Supermarket Sales & Retail Intelligence System")
        self.assertEqual(ws_sum["B5"].value, 3)  # Total records

        # Check Sales Transactions row count (header + 3 rows = 4)
        ws_tx = wb["Sales Transactions"]
        self.assertEqual(ws_tx.max_row, 4)

    def test_excel_workbook_empty_dataframe(self):
        """Verify Excel workbook generates without crashing when passed empty data."""
        empty_df = pd.DataFrame()
        xlsx_buf = generate_excel_workbook(
            df=empty_df,
            dataset_name="Empty Scope",
            filter_summary=None,
            insights=None,
            anomalies=None,
            inventory_df=None,
        )

        self.assertIsInstance(xlsx_buf, io.BytesIO)
        wb = openpyxl.load_workbook(xlsx_buf)
        self.assertIn("Executive Summary", wb.sheetnames)

    def test_csv_export(self):
        """Verify CSV export utility."""
        csv_str = generate_csv_export(self.sample_df)
        self.assertIn("INV-1001", csv_str)
        self.assertIn("INV-1002", csv_str)
        self.assertIn("Organic Apples", csv_str)

        # Empty df test
        empty_csv = generate_csv_export(pd.DataFrame())
        self.assertEqual(empty_csv, "")


class TestReportDownloadAuditIntegration(unittest.TestCase):
    """Regression test suite for report generation, download auditing, and database isolation.

    Validates:
    1. CSV report generation succeeds.
    2. Excel report generation succeeds.
    3. Valid report-generation audit records are saved.
    4. Repeated report downloads work without unexpected duplicate-key failures.
    5. Report logging does not change stock quantities.
    6. Report logging does not create stock movements.
    7. Failed audit logging does not corrupt the database or block file availability.
    8. Existing stock-receipt functionality continues to work seamlessly.
    """

    def setUp(self):
        """Create an isolated temporary SQLite database for each test."""
        self.test_dir = Path(__file__).resolve().parent / "temp"
        self.test_dir.mkdir(parents=True, exist_ok=True)
        self.db_path = self.test_dir / "test_report_audit.db"
        if self.db_path.exists():
            self.db_path.unlink()
        init_database(self.db_path)
        seed_database_if_empty(db_path=self.db_path)

        self.sample_sales = pd.DataFrame([
            {
                "Invoice ID": f"INV-AUDIT-{i}",
                "Branch": "Branch A",
                "City": "Mumbai",
                "Customer Type": "Member",
                "Gender": "Female",
                "Product": "Basmati Rice 5kg",
                "Product Category": "Food & Beverages",
                "Unit Price": 450.0,
                "Quantity": 2,
                "Sales": 900.0,
                "Date": "2026-06-15",
                "Payment": "Ewallet",
                "Rating": 4.8,
            }
            for i in range(1, 6)
        ])

    def tearDown(self):
        """Clean up test database file."""
        if self.db_path.exists():
            try:
                self.db_path.unlink()
            except Exception:
                pass

    def test_01_pdf_download_and_report_logging(self):
        """1. Verify PDF download generation and report logging without integrity error."""
        pdf_buf = generate_pdf_report(
            df=self.sample_sales,
            dataset_name="Test Store Scope",
            filter_summary={"Branch": "Branch A"},
            insights=[],
            anomalies=[],
            inventory_df=None,
        )
        self.assertIsInstance(pdf_buf, io.BytesIO)
        pdf_bytes = pdf_buf.getvalue()
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))

        # Log with short "PDF" and canonical "PDF_EXECUTIVE"
        rep_id = log_report_generation(
            report_type="PDF",
            dataset_scope="SQLite Database (Persistent)",
            filter_summary_json="Branch: Branch A",
            total_sales_analyzed=4500.0,
            total_transactions=5,
            file_name="Executive_Sales_Report_2026-10-18.pdf",
            file_size_bytes=len(pdf_bytes),
            db_path=self.db_path
        )
        self.assertGreater(rep_id, 0)

        history = get_report_history(limit=5, db_path=self.db_path)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["Type"], "PDF_EXECUTIVE")
        self.assertEqual(history[0]["File Name"], "Executive_Sales_Report_2026-10-18.pdf")
        self.assertEqual(history[0]["Records"], 5)

    def test_02_excel_download_and_report_logging(self):
        """2. Verify Excel download generation and report logging without integrity error."""
        excel_buf = generate_excel_workbook(
            df=self.sample_sales,
            dataset_name="Audit Scope",
            filter_summary={"Branch": "Branch A"},
            insights=None,
            anomalies=None,
            inventory_df=None,
        )
        self.assertIsInstance(excel_buf, io.BytesIO)
        excel_bytes = excel_buf.getvalue()
        self.assertGreater(len(excel_bytes), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Executive Summary", wb.sheetnames)
        self.assertIn("Sales Transactions", wb.sheetnames)

        # Log with short "Excel" and canonical "EXCEL_WORKBOOK"
        rep_id = log_report_generation(
            report_type="Excel",
            dataset_scope="SQLite Database (Persistent)",
            filter_summary_json="Branch: Branch A",
            total_sales_analyzed=4500.0,
            total_transactions=5,
            file_name="Supermarket_Intelligence_Report_2026-10-18.xlsx",
            file_size_bytes=len(excel_bytes),
            db_path=self.db_path
        )
        self.assertGreater(rep_id, 0)

        history = get_report_history(limit=5, db_path=self.db_path)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["Type"], "EXCEL_WORKBOOK")
        self.assertEqual(history[0]["File Name"], "Supermarket_Intelligence_Report_2026-10-18.xlsx")
        self.assertEqual(history[0]["Records"], 5)

    def test_03_csv_download_and_report_logging(self):
        """3. Verify CSV download generation and report logging without integrity error."""
        csv_str = generate_csv_export(self.sample_sales)
        self.assertIsInstance(csv_str, str)
        self.assertIn("Basmati Rice 5kg", csv_str)
        csv_bytes = csv_str.encode("utf-8")

        # Log with short "CSV" and canonical "CSV_DATA"
        rep_id = log_report_generation(
            report_type="CSV",
            dataset_scope="SQLite Database (Persistent)",
            filter_summary_json="Branch: Branch A",
            total_sales_analyzed=4500.0,
            total_transactions=5,
            file_name="Filtered_Transactions_2026-10-18.csv",
            file_size_bytes=len(csv_bytes),
            db_path=self.db_path
        )
        self.assertGreater(rep_id, 0)

        history = get_report_history(limit=5, db_path=self.db_path)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["Type"], "CSV_DATA")
        self.assertEqual(history[0]["File Name"], "Filtered_Transactions_2026-10-18.csv")
        self.assertEqual(history[0]["Records"], 5)

    def test_04_valid_audit_records_saved_with_expected_metadata(self):
        """4. Verify valid report-generation audit records are saved with the expected report type and metadata."""
        log_report_generation(
            report_type="CSV",
            dataset_scope="Persistent Scope",
            filter_summary_json="Scope=CSV",
            total_sales_analyzed=1200.5,
            total_transactions=12,
            file_name="data_export.csv",
            file_size_bytes=2048,
            db_path=self.db_path
        )
        log_report_generation(
            report_type="Excel",
            dataset_scope="Persistent Scope",
            filter_summary_json="Scope=Excel",
            total_sales_analyzed=3400.0,
            total_transactions=34,
            file_name="data_workbook.xlsx",
            file_size_bytes=40960,
            db_path=self.db_path
        )
        log_report_generation(
            report_type="PDF",
            dataset_scope="Persistent Scope",
            filter_summary_json="Scope=PDF",
            total_sales_analyzed=5600.0,
            total_transactions=56,
            file_name="briefing.pdf",
            file_size_bytes=81920,
            db_path=self.db_path
        )

        history = get_report_history(limit=10, db_path=self.db_path)
        self.assertEqual(len(history), 3)

        pdf_rec = next((r for r in history if r["Type"] == "PDF_EXECUTIVE"), None)
        excel_rec = next((r for r in history if r["Type"] == "EXCEL_WORKBOOK"), None)
        csv_rec = next((r for r in history if r["Type"] == "CSV_DATA"), None)

        self.assertIsNotNone(pdf_rec)
        self.assertEqual(pdf_rec["File Name"], "briefing.pdf")
        self.assertEqual(pdf_rec["Records"], 56)
        self.assertIn("80.0 KB", str(pdf_rec["File Size"]))

        self.assertIsNotNone(excel_rec)
        self.assertEqual(excel_rec["File Name"], "data_workbook.xlsx")
        self.assertEqual(excel_rec["Records"], 34)
        self.assertIn("40.0 KB", str(excel_rec["File Size"]))

        self.assertIsNotNone(csv_rec)
        self.assertEqual(csv_rec["File Name"], "data_export.csv")
        self.assertEqual(csv_rec["Records"], 12)
        self.assertIn("2.0 KB", str(csv_rec["File Size"]))

    def test_normalize_report_type_aliases_and_rejections(self):
        """Verify normalize_report_type accepts explicit aliases, handles whitespace/casing, and rejects invalid strings."""
        # 1. Valid explicit aliases (exact matches)
        self.assertEqual(normalize_report_type("PDF"), "PDF_EXECUTIVE")
        self.assertEqual(normalize_report_type("PDF_EXECUTIVE"), "PDF_EXECUTIVE")
        self.assertEqual(normalize_report_type("EXCEL"), "EXCEL_WORKBOOK")
        self.assertEqual(normalize_report_type("XLSX"), "EXCEL_WORKBOOK")
        self.assertEqual(normalize_report_type("EXCEL_WORKBOOK"), "EXCEL_WORKBOOK")
        self.assertEqual(normalize_report_type("CSV"), "CSV_DATA")
        self.assertEqual(normalize_report_type("CSV_DATA"), "CSV_DATA")

        # 2. Whitespace and case variations
        self.assertEqual(normalize_report_type("  pdf  "), "PDF_EXECUTIVE")
        self.assertEqual(normalize_report_type("Pdf"), "PDF_EXECUTIVE")
        self.assertEqual(normalize_report_type("excel "), "EXCEL_WORKBOOK")
        self.assertEqual(normalize_report_type("  Xlsx  "), "EXCEL_WORKBOOK")
        self.assertEqual(normalize_report_type("csv"), "CSV_DATA")
        self.assertEqual(normalize_report_type("  CSV_DATA  "), "CSV_DATA")
        self.assertEqual(normalize_report_type("pdf_executive"), "PDF_EXECUTIVE")
        self.assertEqual(normalize_report_type("Excel_Workbook"), "EXCEL_WORKBOOK")

        # 3. Invalid strings (including substrings that must NOT be accepted)
        invalid_types = [
            "NOT_A_PDF_FILE",
            "NOT_CSV",
            "WORD",
            "JSON",
            "HTML",
            "CSV_EXTENDED",
            "MY_EXCEL_FILE",
            "PDF_REPORT",
            "",
            "   ",
            None,
        ]
        for inv in invalid_types:
            with self.assertRaises(ValueError, msg=f"Expected ValueError for '{inv}'"):
                normalize_report_type(inv)

    def test_unsupported_report_types_are_rejected_safely(self):
        """5. Verify unsupported report types are rejected safely with clear ValueError and no database write."""
        history_before = get_report_history(db_path=self.db_path)

        # Test invalid strings specifically requested by user
        for inv in ["NOT_A_PDF_FILE", "NOT_CSV", "WORD", "INVALID_REPORT_TYPE"]:
            with self.assertRaises(ValueError) as ctx:
                log_report(
                    report_type=inv,
                    file_name="test.doc",
                    file_size_bytes=100,
                    dataset_scope="Test Scope",
                    db_path=self.db_path
                )
            self.assertIn("Unsupported report_type", str(ctx.exception))

            with self.assertRaises(ValueError):
                log_report_generation(
                    report_type=inv,
                    dataset_scope="Test Scope",
                    db_path=self.db_path
                )

        # Test empty string and None
        with self.assertRaises(ValueError):
            log_report_generation(report_type="", db_path=self.db_path)

        # Verify no records were inserted
        history_after = get_report_history(db_path=self.db_path)
        self.assertEqual(len(history_before), len(history_after))

    def test_05_repeated_downloads_without_duplicate_key_errors(self):
        """5. Verify repeated report downloads work consecutively without unexpected duplicate-key errors."""
        ids = []
        for i in range(10):
            rid = log_report_generation(
                report_type="CSV_DATA",
                dataset_scope="Persistent Scope",
                filter_summary_json="Test Filter",
                total_sales_analyzed=500.0,
                total_transactions=2,
                file_name="Filtered_Transactions_2026-10-18.csv",  # Same file name repeatedly
                file_size_bytes=512,
                db_path=self.db_path
            )
            self.assertGreater(rid, 0)
            ids.append(rid)

        # Confirm all 10 were inserted with unique sequential auto-increment IDs
        self.assertEqual(len(ids), 10)
        self.assertEqual(len(set(ids)), 10)

    def test_06_report_generation_does_not_alter_stock_quantities(self):
        """6. Verify report generation and logging does not modify inventory stock quantities."""
        prods_before = get_products_df(active_only=False, db_path=self.db_path)
        stocks_before = dict(zip(prods_before["Product ID"], prods_before["Current Stock"]))

        # Execute report generation and logging for all report types
        csv_str = generate_csv_export(self.sample_sales)
        log_report_generation(report_type="CSV", file_size_bytes=len(csv_str), db_path=self.db_path)
        log_report_generation(report_type="Excel", file_size_bytes=5000, db_path=self.db_path)
        log_report_generation(report_type="PDF", file_size_bytes=10000, db_path=self.db_path)

        prods_after = get_products_df(active_only=False, db_path=self.db_path)
        stocks_after = dict(zip(prods_after["Product ID"], prods_after["Current Stock"]))

        self.assertEqual(stocks_before, stocks_after)

    def test_07_report_generation_does_not_create_stock_movements(self):
        """7. Verify report generation and logging does not create inventory movements or alter the ledger."""
        movements_before = get_inventory_movements_df(db_path=self.db_path)
        count_before = len(movements_before)

        # Execute multiple report logs
        log_report_generation(report_type="CSV", db_path=self.db_path)
        log_report_generation(report_type="Excel", db_path=self.db_path)
        log_report_generation(report_type="PDF", db_path=self.db_path)

        movements_after = get_inventory_movements_df(db_path=self.db_path)
        count_after = len(movements_after)

        self.assertEqual(count_before, count_after, "Inventory movements count must remain strictly unchanged.")

    def test_08_simulated_audit_logging_failure_does_not_corrupt_database(self):
        """8. Verify simulated audit-logging failure does not corrupt the database or block an already generated file."""
        # 1. Generate downloadable file first
        csv_str = generate_csv_export(self.sample_sales)
        self.assertGreater(len(csv_str), 0)

        # 2. Simulate logging attempt against an unopenable path (directory path)
        invalid_db_path = self.test_dir
        with self.assertRaises(Exception):
            log_report(
                report_type="CSV_DATA",
                file_name="test.csv",
                file_size_bytes=100,
                dataset_scope="Test",
                db_path=invalid_db_path
            )

        # 3. Verify original database remains completely intact and operational
        valid_prods = get_all_products(db_path=self.db_path)
        self.assertGreater(len(valid_prods), 0)

        # 4. Verify generated file data remains fully intact and usable
        self.assertIn("INV-AUDIT-1", csv_str)

    def test_09_inventory_stock_receipts_continue_to_work(self):
        """9. Verify inventory stock receipts continue to work flawlessly before and after report downloads."""
        all_prods = get_all_products(db_path=self.db_path)
        test_prod = all_prods[0]
        pid = test_prod["product_id"]
        initial_stock = test_prod["current_stock"]

        # 1. Stock receipt (+20 units)
        res1 = record_stock_movement(
            product_id=pid,
            movement_type="RECEIPT",
            quantity_change=20,
            reference_note="Receipt before report download",
            db_path=self.db_path
        )
        self.assertEqual(res1["new_stock_level"], initial_stock + 20)

        # 2. Download and log CSV report
        csv_str = generate_csv_export(self.sample_sales)
        log_report_generation(
            report_type="CSV",
            file_name="Filtered_Transactions.csv",
            file_size_bytes=len(csv_str),
            db_path=self.db_path
        )

        # 3. Stock receipt (+15 units) after report download
        res2 = record_stock_movement(
            product_id=pid,
            movement_type="RECEIPT",
            quantity_change=15,
            reference_note="Receipt after report download",
            db_path=self.db_path
        )
        self.assertEqual(res2["new_stock_level"], initial_stock + 35)

        # 4. Verify final product stock reflects both receipts accurately
        updated_prods = get_all_products(db_path=self.db_path)
        final_prod = next(p for p in updated_prods if p["product_id"] == pid)
        self.assertEqual(final_prod["current_stock"], initial_stock + 35)

    def test_10_existing_database_records_remain_intact(self):
        """10. Verify existing database records (products, categories, sales, anomalies) remain completely intact."""
        # Check initial state
        stats_orig = get_database_stats(db_path=self.db_path)
        prods_orig = get_all_products(db_path=self.db_path)

        # Perform extensive report operations across various formats
        for rtype in ["PDF", "Excel", "CSV", "PDF_EXECUTIVE", "EXCEL_WORKBOOK", "CSV_DATA"]:
            log_report_generation(
                report_type=rtype,
                dataset_scope="Test Scope",
                filter_summary_json="{}",
                total_sales_analyzed=1000.0,
                total_transactions=10,
                file_size_bytes=5000,
                db_path=self.db_path
            )

        # Verify integrity of all other tables
        stats_final = get_database_stats(db_path=self.db_path)
        prods_final = get_all_products(db_path=self.db_path)

        for table in ["categories", "products", "inventory_stock", "inventory_movements", "sales_transactions", "anomaly_records"]:
            self.assertEqual(
                stats_orig[table], stats_final[table],
                f"Table {table} row count changed from {stats_orig[table]} to {stats_final[table]}"
            )

        self.assertEqual(len(prods_orig), len(prods_final))
        self.assertEqual(
            [p["product_id"] for p in prods_orig],
            [p["product_id"] for p in prods_final]
        )
        self.assertEqual(
            [p["current_stock"] for p in prods_orig],
            [p["current_stock"] for p in prods_final]
        )


if __name__ == "__main__":
    unittest.main()
